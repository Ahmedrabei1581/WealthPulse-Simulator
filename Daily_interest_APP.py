import streamlit as st
import pandas as pd
import matplotlib.pyplot as plt
from io import BytesIO
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.pdfgen import canvas
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# --- Page Configuration ---
st.set_page_config(page_title="WMBAP - Smart Banking Return Simulator", page_icon="🏛️", layout="centered")

# --- Language Selector (Sidebar) ---
st.sidebar.title("🌐 Language / اللغة")
lang = st.sidebar.radio("Select Language / اختر اللغة:", ["English", "العربية"], index=0)

# --- Translation Dictionary ---
T = {
    "English": {
        "title_acronym": "🏛️ WMBAP",
        "title_full": "Wealth Management & Banking Analytics Platform",
        "page_desc": "Comprehensive system for simulating daily compounding returns, structured deposits, and investment certificates with executive reports.",
        "client_section": "💼 Client Profile & Credit Information",
        "client_name": "👤 Client / Institutional Entity Name:",
        "account_num": "🔢 Account / Contract Reference:",
        "client_placeholder": "Enter client or institution name...",
        "account_placeholder": "Enter account or contract reference...",
        "engine_select": "🎯 Select Calculation Engine / Banking Product:",
        "products": [
            "📊 Daily Compounding Return (%) - Core Engine",
            "💵 Fixed Daily Amount Simulation",
            "🏦 Time Of Deposit (Periodic Payout)",
            "📜 Certificate Of Deposit 3 Years"
        ],
        "sub_compounding": "📊 Daily Compounding & Effective Annual Rate Analytics",
        "sub_fixed": "💵 Fixed Daily Cash Flow Simulation",
        "sub_deposit": "🏦 Time Of Deposit & Periodic Yield Structure",
        "sub_cert": "📜 3-Year Certificate Of Deposit Management",
        "principal": "Initial Portfolio / Principal Amount:",
        "nominal_rate": "Nominal Annual Rate (%):",
        "daily_amount": "Fixed Daily Payout Amount:",
        "period_days": "Simulation Period (Days):",
        "deposit_months": "Total Deposit Term (Months):",
        "payout_freq": "Interest Payout Frequency:",
        "cert_amount": "Certificate Face Value:",
        "cert_rate": "Fixed Annual Coupon Rate (%):",
        "final_balance": "Final Portfolio Balance after",
        "total_profit": "Total Net Capital Profit:",
        "ear_msg": "Effective Annual Rate (EAR):",
        "simple_rate_msg": "instead of nominal rate",
        "chart_title": "📈 Visual Portfolio Growth Analytics",
        "btn_excel": "📊 Download Executive Excel Report",
        "btn_pdf": "📄 Download Official PDF Certificate",
        "freq_opts": ["Monthly", "Semi-Annual", "Annual"],
        "cert_freq_opts": ["Monthly", "Quarterly", "Semi-Annual", "Annual"],
        "days_label": "Days",
        "years_list": ["Year 1", "Year 2", "Year 3"],
    },
    "العربية": {
        "title_acronym": "🏛️ WMBAP",
        "title_full": "منصة إدارة الثروات والتحليل المالي المصرفي",
        "page_desc": "النظام المتكامل لمحاكاة العوائد التراكمية، الودائع الهيكلية، والشهادات الاستثمارية مع تقارير مؤسسية معتمدة.",
        "client_section": "💼 ملف العميل والبيانات الائتمانية",
        "client_name": "👤 اسم العميل / الجهة الاستثمارية:",
        "account_num": "🔢 رقم الحساب المصرفي / مرجع العقد:",
        "client_placeholder": "أدخل اسم العميل أو الشركة...",
        "account_placeholder": "أدخل رقم الحساب أو المرجع...",
        "engine_select": "🎯 اختر محرك الحساب / المنتج المصرفي المطلوب:",
        "products": [
            "📊 العائد التراكمي اليومي (Daily Compounding %) - المحرك الأساسي",
            "💵 محاكاة العائد الثابت اليومي (Fixed Daily Amount)",
            "🏦 الودائع المصرفية ذات العائد الدوري (Time Of Deposit)",
            "📜 شهادات الاستثمار الثلاثية (Certificate Of Deposit 3 Years)"
        ],
        "sub_compounding": "📊 لوحة تحليل العائد التراكمي وعوائد النمو المركب",
        "sub_fixed": "💵 محاكاة العوائد النقدية الثابتة والمقطوعة يومياً",
        "sub_deposit": "🏦 هيكلة الودائع المصرفية ودوريات صرف العائد",
        "sub_cert": "📜 إدارة الشهادات الادخارية الثلاثية العائد",
        "principal": "أصل المحفظة الاستثمارية / المبلغ المودع:",
        "nominal_rate": "معدل العائد السنوي الاسمي (%):",
        "daily_amount": "قيمة العائد النقدي الثابت اليومي:",
        "period_days": "مدى الفترة الزمنية للمحاكاة (بالأيام):",
        "deposit_months": "إجمالي مدة الوديعة (بالأشهر):",
        "payout_freq": "دورية صرف العائد المصرفي:",
        "cert_amount": "قيمة الشهادة الادخارية:",
        "cert_rate": "معدل العائد الثابت السنوي (%):",
        "final_balance": "إجمالي رصيد المحفظة النهائي بعد",
        "total_profit": "إجمالي صافي الأرباح الرأسمالية:",
        "ear_msg": "معدل العائد السنوي الفعلي (EAR):",
        "simple_rate_msg": "بدلاً من المعدل البسيط",
        "chart_title": "📈 التحليل البصري لتطور نمو المحفظة والأرباح المتراكمة",
        "btn_excel": "📊 تحميل التقرير التنفيذي Excel",
        "btn_pdf": "📄 تحميل الشهادة المعتمدة PDF",
        "freq_opts": ["شهري (Monthly)", "نصف سنوي (Semi-Annual)", "سنوي (Annual)"],
        "cert_freq_opts": ["شهري (Monthly)", "كل 3 شهور (Quarterly)", "نصف سنوي (Semi-Annual)", "سنوي (Annual)"],
        "days_label": "الأيام",
        "years_list": ["السنة الأولى", "السنة الثانية", "السنة الثالثة"],
    }
}

txt = T[lang]

# --- UI Header ---
st.markdown(f"""
    <div style="background: linear-gradient(135deg, #1f4e78 0%, #002060 100%); padding: 25px; border-radius: 12px; color: white; text-align: center; margin-bottom: 25px;">
        <h1 style="margin: 0; font-size: 32px; font-weight: 800; letter-spacing: 1px;">{txt['title_acronym']}</h1>
        <h3 style="margin: 6px 0 0 0; font-size: 18px; font-weight: 500; opacity: 0.95;">{txt['title_full']}</h3>
        <p style="margin: 10px 0 0 0; font-size: 13px; opacity: 0.8; font-style: italic;">{txt['page_desc']}</p>
    </div>
""", unsafe_allow_html=True)

# --- Client Info Section ---
st.markdown(f"### {txt['client_section']}")
with st.container():
    st.markdown('<div style="background-color: #f8f9fa; padding: 15px; border-radius: 8px; border: 1px solid #e5e7eb; margin-bottom: 20px;">', unsafe_allow_html=True)
    col1, col2 = st.columns(2)
    with col1:
        client_name = st.text_input(txt["client_name"], value="", placeholder=txt["client_placeholder"])
    with col2:
        account_number = st.text_input(txt["account_num"], value="", placeholder=txt["account_placeholder"])
    st.markdown('</div>', unsafe_allow_html=True)

# Set defaults if client fields are left empty
display_client = client_name.strip() if client_name.strip() else "N/A"
display_account = account_number.strip() if account_number.strip() else "N/A"

# Product Selection
calculation_type = st.selectbox(txt["engine_select"], txt["products"])
st.markdown("---")

# --- Excel Report Generator ---
def generate_excel_report(client, account, prod_type, summary_dict, df_details):
    output = BytesIO()
    wb = openpyxl.Workbook()
    
    ws_summary = wb.active
    ws_summary.title = "Executive Summary"
    ws_summary.views.sheetView[0].rightToLeft = (lang == "العربية")
    
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=14, bold=True, color="1F4E78")
    border_thin = Border(left=Side(style='thin', color='D3D3D3'), right=Side(style='thin', color='D3D3D3'),
                         top=Side(style='thin', color='D3D3D3'), bottom=Side(style='thin', color='D3D3D3'))
    
    ws_summary.append(["Banking Financial Statement / التقرير المالي المصرفي"])
    ws_summary["A1"].font = title_font
    ws_summary.append([])
    
    ws_summary.append(["KPI / Parameter", "Value"])
    for col in range(1, 3):
        cell = ws_summary.cell(row=3, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        
    row_idx = 4
    for k, v in summary_dict.items():
        ws_summary.append([k, v])
        ws_summary.cell(row=row_idx, column=1).border = border_thin
        ws_summary.cell(row=row_idx, column=2).border = border_thin
        ws_summary.cell(row=row_idx, column=1).alignment = Alignment(horizontal="right")
        ws_summary.cell(row=row_idx, column=2).alignment = Alignment(horizontal="left")
        row_idx += 1
        
    ws_details = wb.create_sheet(title="Detailed Schedule")
    ws_details.views.sheetView[0].rightToLeft = (lang == "العربية")
    
    headers = list(df_details.columns)
    ws_details.append(headers)
    for col in range(1, len(headers) + 1):
        cell = ws_details.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        
    for r_idx, row in enumerate(df_details.values, start=2):
        ws_details.append(list(row))
        for c_idx in range(1, len(row) + 1):
            cell = ws_details.cell(row=r_idx, column=c_idx)
            cell.border = border_thin
            cell.alignment = Alignment(horizontal="center")
            
    wb.save(output)
    output.seek(0)
    return output.getvalue()

# --- PDF Report Generator ---
class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_number(num_pages)
            super().showPage()
        super().save()

    def draw_page_number(self, page_count):
        self.saveState()
        self.setFont("Helvetica", 9)
        self.setFillColor(colors.gray)
        page_text = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(letter[0] - 54, 36, page_text)
        self.drawString(54, 36, "Official Wealth Management & Banking Report (WMBAP)")
        self.restoreState()

def generate_pdf_report(client, account, prod_type_en, summary_data, fig_buffer):
    pdf_buffer = BytesIO()
    doc = SimpleDocTemplate(pdf_buffer, pagesize=letter, rightMargin=54, leftMargin=54, topMargin=54, bottomMargin=54)
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle('ReportTitle', parent=styles['Heading1'], fontName='Helvetica-Bold', fontSize=16, textColor=colors.HexColor('#1f4e78'), alignment=1, spaceAfter=15)
    header_style = ParagraphStyle('ReportHeader', parent=styles['Heading2'], fontName='Helvetica-Bold', fontSize=12, textColor=colors.HexColor('#2c3e50'), spaceAfter=8)
    normal_style = ParagraphStyle('ReportNormal', parent=styles['Normal'], fontName='Helvetica', fontSize=10, textColor=colors.HexColor('#333333'), spaceAfter=6)

    story = []
    story.append(Paragraph("OFFICIAL FINANCIAL & INVESTMENT STATEMENT", title_style))
    story.append(Spacer(1, 10))
    
    client_info = [
        [Paragraph(f"<b>Client Name:</b> {client}", normal_style), Paragraph(f"<b>Account / Contract:</b> {account}", normal_style)],
        [Paragraph(f"<b>Product / Service:</b> {prod_type_en}", normal_style), Paragraph(f"<b>Statement Date:</b> {pd.Timestamp.now().strftime('%Y-%m-%d')}", normal_style)]
    ]
    t_client = Table(client_info, colWidths=[250, 250])
    t_client.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#f2f4f7')),
        ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor('#d1d5db')),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('PADDING', (0, 0), (-1, -1), 8),
    ]))
    story.append(t_client)
    story.append(Spacer(1, 15))
    
    story.append(Paragraph("Executive Financial Summary:", header_style))
    summary_table_data = [[Paragraph(f"<b>{k}</b>", normal_style), Paragraph(str(v), normal_style)] for k, v in summary_data.items()]
    t_summary = Table(summary_table_data, colWidths=[220, 280])
    t_summary.setStyle(TableStyle([
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e5e7eb')),
        ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#f9fafb')),
        ('PADDING', (0, 0), (-1, -1), 6),
    ]))
    story.append(t_summary)
    story.append(Spacer(1, 15))
    
    if fig_buffer:
        story.append(Paragraph("Portfolio Growth & Performance Analytics:", header_style))
        img = Image(fig_buffer, width=420, height=210)
        story.append(img)
        
    doc.build(story, canvasmaker=NumberedCanvas)
    pdf_buffer.seek(0)
    return pdf_buffer.getvalue()

# --- Logic 1: Daily Compounding ---
if calculation_type in [txt["products"][0]]:
    st.subheader(txt["sub_compounding"])
    
    col_c1, col_c2 = st.columns(2)
    with col_c1:
        principal = st.number_input(txt["principal"], value=10000.0, format="%.2f")
    with col_c2:
        annual_rate_percent = st.number_input(txt["nominal_rate"], value=17.85, format="%.2f")
        
    days = st.slider(txt["period_days"], min_value=30, max_value=365, value=365)
    
    daily_rate = (annual_rate_percent / 100.0) / 365.0
    days_list = list(range(1, days + 1))
    balances = [principal * ((1 + daily_rate) ** d) for d in days_list]
    profits = [b - principal for b in balances]
    
    df_raw = pd.DataFrame({
        "Day": days_list,
        "Total Balance": balances,
        "Net Profit": profits
    })
    
    final_amount = balances[-1]
    total_profit = profits[-1]
    effective_rate = ((1 + daily_rate) ** 365 - 1) * 100
    
    st.success(f"**{txt['final_balance']} {days} {txt['days_label']}:** {final_amount:,.2f}")
    st.info(f"**{txt['total_profit']}** {total_profit:,.2f}")
    
    if lang == "العربية":
        st.markdown(f"""
            <div dir="rtl" style="text-align: right; background-color: #d1ecf1; color: #0c5460; padding: 10px 15px; border-radius: 5px; border: 1px solid #bee5eb; font-weight: bold;">
                {txt['ear_msg']} <span dir="ltr" style="unicode-bidi: embed; display: inline-block;">{effective_rate:.2f}%</span> 
                ({txt['simple_rate_msg']} <span dir="ltr" style="unicode-bidi: embed; display: inline-block;">{annual_rate_percent}%</span>)
            </div>
        """, unsafe_allow_html=True)
    else:
        st.warning(f"**{txt['ear_msg']}** {effective_rate:.2f}% ({txt['simple_rate_msg']} {annual_rate_percent}%)")
    
    fig, ax = plt.subplots(figsize=(8, 4))
    ax.plot(days_list, balances, color='#2ca02c', linewidth=2.5, label='Total Portfolio Balance')
    ax.plot(days_list, profits, color='#1f77b4', linestyle='--', linewidth=2, label='Net Capital Profit')
    ax.set_title("Daily Compounding Portfolio Growth Analytics", fontsize=12, fontweight='bold')
    ax.set_xlabel("Days")
    ax.set_ylabel("Portfolio Value")
    ax.legend()
    ax.grid(True, linestyle='--', alpha=0.6)
    
    fig_buf = BytesIO()
    plt.savefig(fig_buf, format='png', bbox_inches='tight', dpi=150)
    fig_buf.seek(0)
    plt.close(fig)
    
    st.subheader(txt["chart_title"])
    st.line_chart(pd.DataFrame({"Total Balance": balances, "Net Profit": profits}, index=days_list))
    
    summary_dict = {
        "Client Name": display_client,
        "Account / Contract": display_account,
        "Initial Principal": f"{principal:,.2f}",
        "Nominal Annual Rate": f"{annual_rate_percent}%",
        "Effective Annual Rate (EAR)": f"{effective_rate:.2f}%",
        "Simulation Period (Days)": days,
        "Final Portfolio Balance": f"{final_amount:,.2f}",
        "Total Net Capital Profit": f"{total_profit:,.2f}"
    }
    
    col_dl1, col_dl2 = st.columns(2)
    with col_dl1:
        excel_data = generate_excel_report(display_client, display_account, "Daily Compounding %", summary_dict, df_raw)
        st.download_button(txt["btn_excel"], data=excel_data, file_name=f'compounding_return_{display_account}.xlsx', mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    with col_dl2:
        pdf_data = generate_pdf_report(display_client, display_account, "Daily Compounding %", summary_dict, fig_buf)
        st.download_button(txt["btn_pdf"], data=pdf_data, file_name=f'investment_report_{display_account}.pdf', mime='application/pdf')

# --- Logic 2: Fixed Daily Amount ---
elif calculation_type in [txt["products"][1]]:
    st.subheader(txt["sub_fixed"])
    
    col_p1, col_p2 = st.columns(2)
    with col_p1:
        principal = st.number_input(txt["principal"], value=100000.0, format="%.2f")
    with col_p2:
        daily_amount = st.number_input(txt["daily_amount"], value=17.85, format="%.2f")
        
    days = st.slider(txt["period_days"], min_value=30, max_value=365, value=365)
    
    days_list = list(range(1, days + 1))
    cumulative_returns = [daily_amount * d for d in days_list]
    total_balances = [principal + cr for cr in cumulative_returns]
    
    df = pd.DataFrame({
        "Day": days_list,
        "Total Return": cumulative_returns,
        "Total Balance": total_balances
    })
    
    final_return = cumulative_returns[-1]
    final_total = total_balances[-1]
    
    st.success(f"**{txt['total_profit']}** {final_return:,.2f}")
    st.info(f"**{txt['final_balance']} {days} {txt['days_label']}:** {final_total:,.2f}")
    
    fig, ax = plt.subplots(figsize=(8, 4))
    ax.plot(days_list, total_balances, color='#1f4e78', linewidth=2.5, label='Total Account Balance')
    ax.set_title("Fixed Daily Payout Cash Flow Analytics", fontsize=12, fontweight='bold')
    ax.set_xlabel("Days")
    ax.set_ylabel("Total Amount")
    ax.legend()
    ax.grid(True, linestyle='--', alpha=0.6)
    
    fig_buf = BytesIO()
    plt.savefig(fig_buf, format='png', bbox_inches='tight', dpi=150)
    fig_buf.seek(0)
    plt.close(fig)
    
    st.subheader(txt["chart_title"])
    st.line_chart(pd.DataFrame({"Total Balance": total_balances}, index=days_list))
    
    summary_dict = {
        "Client Name": display_client,
        "Account / Contract": display_account,
        "Initial Principal": f"{principal:,.2f}",
        "Daily Fixed Return": f"{daily_amount:,.2f}",
        "Period (Days)": days,
        "Total Accumulated Return": f"{final_return:,.2f}",
        "Final Total Balance": f"{final_total:,.2f}"
    }
    
    col_dl1, col_dl2 = st.columns(2)
    with col_dl1:
        excel_data = generate_excel_report(display_client, display_account, "Fixed Daily Amount", summary_dict, df)
        st.download_button(txt["btn_excel"], data=excel_data, file_name=f'fixed_return_{display_account}.xlsx', mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    with col_dl2:
        pdf_data = generate_pdf_report(display_client, display_account, "Fixed Daily Amount", summary_dict, fig_buf)
        st.download_button(txt["btn_pdf"], data=pdf_data, file_name=f'financial_report_{display_account}.pdf', mime='application/pdf')

# --- Logic 3: Time Of Deposit ---
elif calculation_type in [txt["products"][2]]:
    st.subheader(txt["sub_deposit"])
    
    col_d1, col_d2 = st.columns(2)
    with col_d1:
        deposit_amount = st.number_input(txt["principal"], value=200000.0, format="%.2f")
        deposit_rate = st.number_input(txt["nominal_rate"], value=18.0, format="%.2f")
    with col_d2:
        payout_freq_selected = st.selectbox(txt["payout_freq"], txt["freq_opts"])
        deposit_months = st.slider(txt["deposit_months"], min_value=1, max_value=60, value=12)

    if "Monthly" in payout_freq_selected or "شهري" in payout_freq_selected:
        periods_per_year = 12
        payout_freq_en = "Monthly"
    elif "Semi-Annual" in payout_freq_selected or "نصف" in payout_freq_selected:
        periods_per_year = 2
        payout_freq_en = "Semi-Annual"
    else:
        periods_per_year = 1
        payout_freq_en = "Annual"

    total_periods = int((deposit_months / 12) * periods_per_year)
    periodic_rate = (deposit_rate / 100.0) / periods_per_year
    periodic_profit = deposit_amount * periodic_rate
    total_profit_deposit = periodic_profit * total_periods

    st.success(f"**Profit Per Payout Period / صافي العائد النقدي:** {periodic_profit:,.2f}")
    st.info(f"**{txt['total_profit']}** {total_profit_deposit:,.2f}")
    
    period_list = list(range(1, total_periods + 1))
    profits_cumulative = [periodic_profit * p for p in period_list]
    
    df_dep = pd.DataFrame({
        "Period": period_list,
        "Periodic Profit": [periodic_profit] * total_periods,
        "Cumulative Profit": profits_cumulative
    })
    
    fig, ax = plt.subplots(figsize=(8, 4))
    ax.bar(period_list, [periodic_profit]*len(period_list), color='#2ca02c', width=0.4)
    ax.set_title(f"Time Of Deposit Periodic Payouts ({payout_freq_en})", fontsize=12, fontweight='bold')
    ax.set_xlabel("Payout Period")
    ax.set_ylabel("Profit Amount")
    ax.grid(True, linestyle='--', alpha=0.6)
    
    fig_buf = BytesIO()
    plt.savefig(fig_buf, format='png', bbox_inches='tight', dpi=150)
    fig_buf.seek(0)
    plt.close(fig)
    
    summary_dict = {
        "Client Name": display_client,
        "Account / Contract": display_account,
        "Deposit Principal": f"{deposit_amount:,.2f}",
        "Annual Interest Rate": f"{deposit_rate}%",
        "Payout Frequency": payout_freq_en,
        "Deposit Duration (Months)": deposit_months,
        "Profit Per Payout Period": f"{periodic_profit:,.2f}",
        "Total Accumulated Profit": f"{total_profit_deposit:,.2f}"
    }
    
    col_dl1, col_dl2 = st.columns(2)
    with col_dl1:
        excel_data = generate_excel_report(display_client, display_account, "Time Of Deposit", summary_dict, df_dep)
        st.download_button(txt["btn_excel"], data=excel_data, file_name=f'deposit_report_{display_account}.xlsx', mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    with col_dl2:
        pdf_data = generate_pdf_report(display_client, display_account, "Time Of Deposit", summary_dict, fig_buf)
        st.download_button(txt["btn_pdf"], data=pdf_data, file_name=f'financial_report_{display_account}.pdf', mime='application/pdf')

# --- Logic 4: Certificate Of Deposit 3 Years ---
else:
    st.subheader(txt["sub_cert"])
    
    col_s1, col_s2 = st.columns(2)
    with col_s1:
        cert_amount = st.number_input(txt["cert_amount"], value=500000.0, format="%.2f")
        cert_rate = st.number_input(txt["cert_rate"], value=20.0, format="%.2f")
    with col_s2:
        cert_freq_selected = st.selectbox(txt["payout_freq"], txt["cert_freq_opts"])

    if "Monthly" in cert_freq_selected or "شهري" in cert_freq_selected:
        c_periods = 12
        cert_freq_en = "Monthly"
    elif "Quarterly" in cert_freq_selected or "كل 3" in cert_freq_selected:
        c_periods = 4
        cert_freq_en = "Quarterly"
    elif "Semi-Annual" in cert_freq_selected or "نصف" in cert_freq_selected:
        c_periods = 2
        cert_freq_en = "Semi-Annual"
    else:
        c_periods = 1
        cert_freq_en = "Annual"

    annual_profit_cert = cert_amount * (cert_rate / 100.0)
    periodic_cert_profit = annual_profit_cert / c_periods
    total_3yr_profit = annual_profit_cert * 3

    st.success(f"**Profit Per Payout Period / العائد النقدي:** {periodic_cert_profit:,.2f}")
    st.info(f"**{txt['total_profit']}** {total_3yr_profit:,.2f}")

    years_list = txt["years_list"]
    annual_profits_list = [annual_profit_cert, annual_profit_cert, annual_profit_cert]
    
    df_cert = pd.DataFrame({
        "Year": years_list,
        "Annual Profit": annual_profits_list,
        "Total Profit to Date": [annual_profit_cert, annual_profit_cert*2, annual_profit_cert*3]
    })
    
    fig, ax = plt.subplots(figsize=(8, 4))
    ax.bar(years_list, annual_profits_list, color='#1f77b4', width=0.4)
    ax.set_title("Certificate Of Deposit 3 Years - Annual Profits", fontsize=12, fontweight='bold')
    ax.set_xlabel("Years")
    ax.set_ylabel("Annual Profit Amount")
    ax.grid(True, linestyle='--', alpha=0.6)
    
    fig_buf = BytesIO()
    plt.savefig(fig_buf, format='png', bbox_inches='tight', dpi=150)
    fig_buf.seek(0)
    plt.close(fig)
    
    summary_dict = {
        "Client Name": display_client,
        "Account / Contract": display_account,
        "Certificate Principal": f"{cert_amount:,.2f}",
        "Annual Interest Rate": f"{cert_rate}%",
        "Payout Frequency": cert_freq_en,
        "Investment Duration": "3 Years",
        "Profit Per Payout Period": f"{periodic_cert_profit:,.2f}",
        "Total 3-Year Cumulative Profit": f"{total_3yr_profit:,.2f}"
    }
    
    col_dl1, col_dl2 = st.columns(2)
    with col_dl1:
        excel_data = generate_excel_report(display_client, display_account, "Certificate Of Deposit 3 Years", summary_dict, df_cert)
        st.download_button(txt["btn_excel"], data=excel_data, file_name=f'certificate_report_{display_account}.xlsx', mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    with col_dl2:
        pdf_data = generate_pdf_report(display_client, display_account, "Certificate Of Deposit 3 Years", summary_dict, fig_buf)
        st.download_button(txt["btn_pdf"], data=pdf_data, file_name=f'financial_report_{display_account}.pdf', mime='application/pdf')
